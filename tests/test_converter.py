from pathlib import Path
import sqlite3

import pytest
from openpyxl import Workbook

from production_planner.converter import SHEETS, convert_workbook, read_workbook
from production_planner.database import validate_database
from production_planner.errors import ValidationError


def _workbook(path: Path, cycle: bool = False):
    workbook = Workbook()
    workbook.remove(workbook.active)
    data = {
        "Metadata": [["database_name", "Fixture"]],
        "Items": [["raw", "Raw", "raw", "unit", 1, "test", None], ["final", "Final", "producible", "unit", 5, "test", None]],
        "Stations": [["machine", "Machine", None]],
        "Recipes": [["make_final", "final", "Make", 1, "machine", 10, 1, None]],
        "Ingredients": [["make_final", "final" if cycle else "raw", 1]],
        "CapacityProfiles": [["default", "Default"]],
        "Capacities": [["default", "machine", 1]],
        "Sources": [["fixture", "Fixture", None, None, None]],
    }
    for name, headers in SHEETS.items():
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for row in data[name]:
            sheet.append(row)
    workbook.save(path)


def test_converts_valid_workbook_atomically(tmp_path):
    source, target = tmp_path / "input.xlsx", tmp_path / "output.sqlite"
    _workbook(source)
    convert_workbook(source, target)
    validate_database(target)
    with pytest.raises(ValidationError, match="target already exists"):
        convert_workbook(source, target)
    convert_workbook(source, target, force=True)
    validate_database(target)


def test_rejects_cycle_before_database_creation(tmp_path):
    source, target = tmp_path / "cycle.xlsx", tmp_path / "output.sqlite"
    _workbook(source, cycle=True)
    with pytest.raises(ValidationError, match="cycle"):
        read_workbook(source)
    assert not target.exists()


def test_converts_normalized_database_workbook_and_preserves_extended_fields(tmp_path):
    source, target = tmp_path / "normalized.xlsx", tmp_path / "normalized.sqlite"
    workbook = Workbook()
    schema = workbook.active
    schema.title = "schema"
    schema.append(["metadata_key", "metadata_value"])
    schema.append(["data_as_of", "2026-09-24"])
    sheets = {
        "items": (["item_id", "item_name", "category_id", "unlock_level", "source_url"], [
            ["item_wheat", "Wheat", "cat_crop", 1, "https://example.test/wheat"],
            ["item_bread", "Bread", "cat_good", 2, "https://example.test/bread"],
            ["item_voucher", "Voucher", "cat_reward", None, "https://example.test/voucher"],
        ]),
        "production_methods": (["production_method_id", "output_item_id", "station_id", "output_quantity", "base_time_minutes", "mastered_time_minutes", "unlock_level", "notes"], [
            ["method_wheat", "item_wheat", "station_field", 1, 2, None, 1, None],
            ["method_bread", "item_bread", "station_bakery", 1, 5, 4, 2, None],
        ]),
        "production_components": (["production_method_id", "component_item_id", "component_quantity", "component_sequence"], [
            ["method_bread", "item_wheat", 3, 1], ["method_bread", "item_voucher", 1, 2],
        ]),
        "stations": (["station_id", "station_name", "station_type"], [
            ["station_field", "Field", "field"], ["station_bakery", "Bakery", "production_machine"],
        ]),
        "categories": (["category_id", "category_name", "description"], [
            ["cat_crop", "Crop", "Grown"], ["cat_good", "Good", "Produced"], ["cat_reward", "Reward", "Awarded"],
        ]),
    }
    for name, (headers, data) in sheets.items():
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for row in data:
            sheet.append(row)
    workbook.save(source)

    convert_workbook(source, target)
    with sqlite3.connect(target) as connection:
        assert connection.execute("SELECT COUNT(*) FROM items").fetchone()[0] == 3
        assert connection.execute("SELECT kind FROM items WHERE item_key='item_voucher'").fetchone()[0] == "raw"
        assert connection.execute("SELECT average_value FROM items WHERE item_key='item_bread'").fetchone()[0] == "1"
        assert connection.execute("SELECT mastered_duration_seconds FROM recipes WHERE recipe_key='method_bread'").fetchone()[0] == 240
        assert connection.execute("SELECT station_type FROM stations WHERE station_key='station_bakery'").fetchone()[0] == "production_machine"
        assert connection.execute("SELECT COUNT(*) FROM capacities WHERE station_count=1").fetchone()[0] == 2
