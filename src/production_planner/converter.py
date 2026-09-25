from __future__ import annotations

import math
import os
import tempfile
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from .database import create_database
from .errors import ValidationError

SHEETS = {
    "Metadata": ["key", "value"],
    "Items": ["item_key", "name", "kind", "unit", "average_value", "value_basis", "notes"],
    "Stations": ["station_key", "name", "notes"],
    "Recipes": ["recipe_key", "output_item_key", "process_name", "output_quantity", "station_key", "duration_seconds", "is_active", "notes"],
    "Ingredients": ["recipe_key", "item_key", "quantity"],
    "CapacityProfiles": ["profile_key", "name"],
    "Capacities": ["profile_key", "station_key", "station_count"],
    "Sources": ["source_key", "title", "url", "accessed_date", "notes"],
}

NORMALIZED_SHEETS = {
    "items": ["item_id", "item_name", "category_id", "unlock_level", "source_url"],
    "production_methods": ["production_method_id", "output_item_id", "station_id", "output_quantity", "base_time_minutes", "mastered_time_minutes", "unlock_level", "notes"],
    "production_components": ["production_method_id", "component_item_id", "component_quantity", "component_sequence"],
    "stations": ["station_id", "station_name", "station_type"],
    "categories": ["category_id", "category_name", "description"],
}


def _clean(value: object) -> object | None:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _rows(sheet, headers: list[str]) -> list[dict[str, object | None]]:
    actual = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
    if actual[: len(headers)] != headers:
        raise ValidationError(f"sheet {sheet.title} must begin with columns: {', '.join(headers)}")
    result = []
    for values in sheet.iter_rows(min_row=2, max_col=len(headers), values_only=True):
        if all(_clean(value) is None for value in values):
            continue
        result.append({header: _clean(value) for header, value in zip(headers, values, strict=True)})
    return result


def _decimal(value: object, label: str, positive: bool = True) -> str:
    try:
        number = Decimal(str(value))
    except (InvalidOperation, TypeError) as exc:
        raise ValidationError(f"{label} must be a number") from exc
    if not number.is_finite() or (positive and number <= 0):
        raise ValidationError(f"{label} must be positive")
    return str(number.normalize())


def read_workbook(path: Path) -> dict[str, list[dict[str, object]]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError(f"could not read workbook: {exc}") from exc
    try:
        if set(NORMALIZED_SHEETS) <= set(workbook.sheetnames):
            result = _read_normalized_workbook(workbook)
        else:
            missing = set(SHEETS) - set(workbook.sheetnames)
            if missing:
                raise ValidationError(f"missing sheets: {', '.join(sorted(missing))}")
            rows = {key: _rows(workbook[key], headers) for key, headers in SHEETS.items()}
            result = {
                "metadata": rows["Metadata"],
                "categories": [],
                "items": rows["Items"],
                "stations": rows["Stations"],
                "recipes": rows["Recipes"],
                "ingredients": rows["Ingredients"],
                "capacity_profiles": rows["CapacityProfiles"],
                "capacities": rows["Capacities"],
                "sources": rows["Sources"],
            }
    finally:
        workbook.close()
    _validate(result)
    return result


def _read_normalized_workbook(workbook) -> dict[str, list[dict[str, object]]]:
    rows = {name: _rows(workbook[name], headers) for name, headers in NORMALIZED_SHEETS.items()}
    method_outputs = {str(row["output_item_id"]) for row in rows["production_methods"]}
    station_names = {str(row["station_id"]): str(row["station_name"]) for row in rows["stations"]}

    metadata: list[dict[str, object]] = [
        {"key": "database_name", "value": "Hay Day production catalog"},
        {"key": "source_format", "value": "normalized_hay_day_workbook"},
    ]
    if "schema" in workbook.sheetnames:
        schema_values = list(workbook["schema"].iter_rows(values_only=True))
        start = next((index for index, row in enumerate(schema_values) if row and row[0] == "metadata_key"), None)
        if start is not None:
            for row in schema_values[start + 1:]:
                if not row or row[0] is None:
                    break
                metadata.append({"key": str(row[0]), "value": str(row[1])})

    categories = [{
        "category_key": row["category_id"], "name": row["category_name"], "description": row["description"],
    } for row in rows["categories"]]
    items = [{
        "item_key": row["item_id"], "name": row["item_name"],
        "kind": "producible" if str(row["item_id"]) in method_outputs else "raw",
        "unit": "unit", "average_value": "1", "value_basis": "temporary v1.1 default",
        "notes": None, "category_key": row["category_id"], "unlock_level": row["unlock_level"],
        "source_url": row["source_url"],
    } for row in rows["items"]]
    stations = [{
        "station_key": row["station_id"], "name": row["station_name"], "notes": None,
        "station_type": row["station_type"],
    } for row in rows["stations"]]
    recipes = [{
        "recipe_key": row["production_method_id"], "output_item_key": row["output_item_id"],
        "process_name": station_names.get(str(row["station_id"]), str(row["station_id"])),
        "output_quantity": row["output_quantity"], "station_key": row["station_id"],
        "duration_seconds": int(row["base_time_minutes"] * 60), "is_active": 1, "notes": row["notes"],
        "mastered_duration_seconds": None if row["mastered_time_minutes"] is None else int(row["mastered_time_minutes"] * 60),
        "unlock_level": row["unlock_level"],
    } for row in rows["production_methods"]]
    ingredients = [{
        "recipe_key": row["production_method_id"], "item_key": row["component_item_id"],
        "quantity": row["component_quantity"], "component_sequence": row["component_sequence"],
    } for row in rows["production_components"]]
    return {
        "metadata": metadata,
        "categories": categories,
        "items": items,
        "stations": stations,
        "recipes": recipes,
        "ingredients": ingredients,
        "capacity_profiles": [{"profile_key": "hay_day_default", "name": "Hay Day default"}],
        "capacities": [{"profile_key": "hay_day_default", "station_key": row["station_id"], "station_count": 1} for row in rows["stations"]],
        "sources": [],
    }


def _unique(rows: list[dict[str, object]], field: str, label: str) -> set[str]:
    keys: list[str] = []
    for row in rows:
        value = row.get(field)
        if value is None:
            raise ValidationError(f"{label} has a blank {field}")
        keys.append(str(value))
    duplicates = sorted({key for key in keys if keys.count(key) > 1})
    if duplicates:
        raise ValidationError(f"duplicate {label} keys: {', '.join(duplicates)}")
    return set(keys)


def _validate(rows: dict[str, list[dict[str, object]]]) -> None:
    metadata_keys = _unique(rows["metadata"], "key", "metadata")
    if "schema_version" in metadata_keys:
        raise ValidationError("metadata key schema_version is reserved")
    item_keys = _unique(rows["items"], "item_key", "item")
    category_keys = _unique(rows.get("categories", []), "category_key", "category")
    station_keys = _unique(rows["stations"], "station_key", "station")
    recipe_keys = _unique(rows["recipes"], "recipe_key", "recipe")
    profile_keys = _unique(rows["capacity_profiles"], "profile_key", "capacity profile")
    _unique(rows["sources"], "source_key", "source")

    for row in rows["items"]:
        if not row.get("name"):
            raise ValidationError(f"item {row['item_key']} has a blank name")
        if row.get("kind") not in {"raw", "producible"}:
            raise ValidationError(f"item {row['item_key']} kind must be raw or producible")
        row["unit"] = row.get("unit") or "unit"
        if row.get("average_value") is not None:
            row["average_value"] = _decimal(row["average_value"], f"item {row['item_key']} average_value", positive=False)
        if row.get("category_key") is not None and str(row["category_key"]) not in category_keys:
            raise ValidationError(f"item {row['item_key']} references an unknown category")
    active_by_output: dict[str, int] = defaultdict(int)
    for row in rows["recipes"]:
        if str(row.get("output_item_key")) not in item_keys or str(row.get("station_key")) not in station_keys:
            raise ValidationError(f"recipe {row['recipe_key']} references an unknown item or station")
        if not row.get("process_name"):
            raise ValidationError(f"recipe {row['recipe_key']} has a blank process_name")
        row["output_quantity"] = _decimal(row.get("output_quantity"), f"recipe {row['recipe_key']} output_quantity")
        try:
            row["duration_seconds"] = int(row.get("duration_seconds"))
            row["is_active"] = int(1 if row.get("is_active") is None else row["is_active"])
        except (TypeError, ValueError) as exc:
            raise ValidationError(f"recipe {row['recipe_key']} duration/is_active is invalid") from exc
        if row["duration_seconds"] < 0 or row["is_active"] not in {0, 1}:
            raise ValidationError(f"recipe {row['recipe_key']} duration/is_active is invalid")
        active_by_output[str(row["output_item_key"])] += row["is_active"]
    producible = {str(row["item_key"]) for row in rows["items"] if row["kind"] == "producible"}
    invalid_outputs = sorted(str(row["output_item_key"]) for row in rows["recipes"] if str(row["output_item_key"]) not in producible)
    if invalid_outputs:
        raise ValidationError(f"recipes may output only producible items: {', '.join(invalid_outputs)}")
    if {key for key, count in active_by_output.items() if count > 1}:
        raise ValidationError("an item may have only one active recipe")
    missing_recipes = sorted(key for key in producible if active_by_output[key] != 1)
    if missing_recipes:
        raise ValidationError(f"producible items need one active recipe: {', '.join(missing_recipes)}")
    seen_ingredients: set[tuple[str, str]] = set()
    deps: dict[str, set[str]] = defaultdict(set)
    outputs = {str(row["recipe_key"]): str(row["output_item_key"]) for row in rows["recipes"] if row["is_active"]}
    for row in rows["ingredients"]:
        recipe_key, item_key = str(row.get("recipe_key")), str(row.get("item_key"))
        if recipe_key not in recipe_keys or item_key not in item_keys:
            raise ValidationError("ingredient references an unknown recipe or item")
        if (recipe_key, item_key) in seen_ingredients:
            raise ValidationError(f"duplicate ingredient {item_key} in recipe {recipe_key}")
        seen_ingredients.add((recipe_key, item_key))
        row["quantity"] = _decimal(row.get("quantity"), f"ingredient {recipe_key}/{item_key}")
        if recipe_key in outputs and item_key in producible:
            deps[outputs[recipe_key]].add(item_key)
    visiting: set[str] = set()
    visited: set[str] = set()
    def visit(key: str) -> None:
        if key in visiting:
            raise ValidationError(f"recipe cycle detected at {key}")
        if key in visited:
            return
        visiting.add(key)
        for dependency in deps[key]:
            visit(dependency)
        visiting.remove(key)
        visited.add(key)
    for key in producible:
        visit(key)
    capacity_pairs: set[tuple[str, str]] = set()
    for row in rows["capacity_profiles"]:
        if not row.get("name"):
            raise ValidationError(f"capacity profile {row['profile_key']} has a blank name")
    for row in rows["stations"]:
        if not row.get("name"):
            raise ValidationError(f"station {row['station_key']} has a blank name")
    for row in rows["sources"]:
        if not row.get("title"):
            raise ValidationError(f"source {row['source_key']} has a blank title")
    for row in rows["capacities"]:
        if str(row.get("profile_key")) not in profile_keys or str(row.get("station_key")) not in station_keys:
            raise ValidationError("capacity references an unknown profile or station")
        pair = (str(row["profile_key"]), str(row["station_key"]))
        if pair in capacity_pairs:
            raise ValidationError(f"duplicate capacity for {pair[0]}/{pair[1]}")
        capacity_pairs.add(pair)
        try:
            row["station_count"] = int(row.get("station_count"))
        except (TypeError, ValueError) as exc:
            raise ValidationError("station_count must be an integer") from exc
        if row["station_count"] < 1:
            raise ValidationError("station_count must be at least 1")


def convert_workbook(source: Path, target: Path, force: bool = False) -> Path:
    if target.exists() and not force:
        raise ValidationError(f"target already exists: {target}; pass --force to replace it")
    rows = read_workbook(source)
    target.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_name = tempfile.mkstemp(prefix="planner-", suffix=".sqlite", dir=target.parent)
    os.close(fd)
    temporary = Path(temporary_name)
    temporary.unlink()
    try:
        create_database(temporary, rows)
        temporary.replace(target)
    finally:
        if temporary.exists():
            temporary.unlink()
    return target
